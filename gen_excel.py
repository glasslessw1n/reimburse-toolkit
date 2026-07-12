#!/usr/bin/env python3
"""
生成报销明细 Excel：按差旅区间逐项列出
- 铁路交通：日期、金额、起始站
- 酒店：入住-退房日期、金额、酒店城市
- 滴滴出行：区间总额、返回城市及日期
- 餐饮：日期、金额
- 通信：月份、金额

用法:
  python3 gen_excel.py --base-dir /path/to/bills
  python3 gen_excel.py --base-dir . --output 我的报销明细.xlsx
"""
import sys, re, subprocess
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent))
from organize import discover_trips, classify_and_sort, classify_local, parse_date_range, get_base_dir

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# 解析命令行参数
BASE_DIR = get_base_dir()
OUTPUT = None
i = 1
while i < len(sys.argv):
    arg = sys.argv[i]
    if arg == '--output' and i + 1 < len(sys.argv):
        OUTPUT = Path(sys.argv[i + 1])
        i += 2
    elif arg.startswith('--output='):
        OUTPUT = Path(arg.split('=', 1)[1])
        i += 1
    else:
        i += 1

if OUTPUT is None:
    trip_dirs_preview, _ = discover_trips(BASE_DIR)
    sorted_dirs = sorted(trip_dirs_preview)
    if sorted_dirs:
        first = re.findall(r'(\d{4})', sorted_dirs[0])
        last = re.findall(r'(\d{4})', sorted_dirs[-1])
        if first and last:
            s, e = first[0], last[-1]
            OUTPUT = BASE_DIR / f"报销明细_{s[:2]}{s[2:]}-{e[:2]}{e[2:]}.xlsx"
        else:
            OUTPUT = BASE_DIR / "报销明细.xlsx"
    else:
        OUTPUT = BASE_DIR / "报销明细.xlsx"


def extract_pdf_amount(filepath):
    """从 PDF 提取金额：优先取价税合计（小写）"""
    try:
        result = subprocess.run(
            ['pdftotext', str(filepath), '-'],
            capture_output=True, text=True, timeout=5
        )
        text = result.stdout
        compact = re.sub(r'\s+', '', text)

        # 策略1：标准 小写¥金额 模式（适用于大多数发票）
        m = re.search(r'[（(]小写[）)]\s*[¥￥]([\d,]+\.?\d{0,2})', compact)
        if m:
            val = float(m.group(1).replace(',', ''))
            if val > 10:  # 正常金额，直接返回
                return val
            # val ≤ 10 说明金额被PDF排版截断，需要拼接修复
            idx = compact.find('小写')
            if idx > 0:
                before = compact[max(0, idx-40):idx]
                after = compact[idx:idx+30]
                before_nums = re.findall(r'(\d+\.\d{1,2})', before)
                after_match = re.search(r'[¥￥]([\d]+)', after)
                if before_nums and after_match:
                    suffix = before_nums[-1]
                    prefix = after_match.group(1)
                    try:
                        combined = float(prefix + suffix)
                        if 10 < combined < 100000:
                            return combined
                    except:
                        pass

        # 策略2：价税合计
        m = re.search(r'价税合计.*?[¥￥]([\d,]+\.?\d{0,2})', compact)
        if m:
            return float(m.group(1).replace(',', ''))

        # 策略3：最后一个¥金额（取最大值更安全）
        amounts = re.findall(r'[¥￥]([\d,]+\.?\d{0,2})', compact)
        if amounts:
            vals = [float(a.replace(',', '')) for a in amounts]
            vals.sort()
            return vals[-1]  # 取最大金额
    except:
        pass
    return None


# ── 样式 ──────────────────────────────────────
header_font = Font(name="微软雅黑", bold=True, size=11, color="FFFFFF")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
cell_align = Alignment(vertical="center", wrap_text=True)
center_align = Alignment(horizontal="center", vertical="center")
right_align = Alignment(horizontal="right", vertical="center")
thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)
cat_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
subtotal_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
trip_total_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
grand_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
amount_fmt = '#,##0.00'

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "报销明细"

# 表头
headers = ["票据类别", "日期/区间", "金额(¥)", "明细", "备注"]
widths = [14, 18, 12, 30, 24]
for col, (h, w) in enumerate(zip(headers, widths), 1):
    c = ws.cell(row=1, column=col, value=h)
    c.font = header_font
    c.fill = header_fill
    c.alignment = header_align
    c.border = thin_border
    ws.column_dimensions[get_column_letter(col)].width = w
ws.freeze_panes = "A2"

def w(row, col, value, font=None, fill=None, align=None, fmt=None):
    """写单元格（跳过合并区域）"""
    try:
        c = ws.cell(row=row, column=col)
        c.value = value
        c.font = font or Font(name="微软雅黑", size=10)
        c.border = thin_border
        c.alignment = align or cell_align
        if fill:
            c.fill = fill
        if fmt:
            c.number_format = fmt
        return c
    except AttributeError:
        return None

def merge_and_write(row, col1, col2, value, **kwargs):
    """合并单元格并写入"""
    ws.merge_cells(start_row=row, start_column=col1, end_row=row, end_column=col2)
    return w(row, col1, value, **kwargs)


# ── 主循环 ────────────────────────────────────
trip_dirs, _ = discover_trips(BASE_DIR)
row = 2
totals = {"交通": 0, "酒店": 0, "滴滴": 0, "餐饮": 0, "通信": 0}

for trip_dir in trip_dirs:
    data = classify_and_sort(trip_dir, BASE_DIR)
    if not data:
        continue

    # ── 行程标题行 ──
    merge_and_write(row, 1, 5, f"【{trip_dir}】",
                    font=Font(name="微软雅黑", bold=True, size=12), fill=cat_fill)
    row += 1

    # ── 1. 铁路交通 ──
    transport = sorted(data["transport"],
                       key=lambda t: re.search(r"(\d{4})", t["filename"]).group(1) if re.search(r"(\d{4})", t["filename"]) else "9999")

    trip_t = 0
    w(row, 1, "铁路交通", font=Font(name="微软雅黑", bold=True, size=10))

    for t in transport:
        name = t["filename"]
        row += 1

        date_m = re.search(r"(\d{4})", name)
        mm, dd = (int(date_m.group(1)[:2]), date_m.group(1)[2:]) if date_m else ("", "")
        date_str = f"{mm:02d}/{dd}" if date_m else ""

        is_boarding = "登机牌" in name
        if not is_boarding:
            amt_m = re.search(r"(\d+\.?\d*)\.pdf", name)
            amt = float(amt_m.group(1)) if amt_m else None
        else:
            amt = None

        if amt:
            trip_t += amt

        route_m = re.search(r"(\S+-\S+)", name)
        route = route_m.group(1) if route_m else ""

        w(row, 1, "登机牌" if is_boarding else "火车/机票")
        w(row, 2, date_str, align=center_align)
        w(row, 3, amt if amt else "-", align=right_align, fmt=amount_fmt if amt else None)
        w(row, 4, route)
        w(row, 5, name)

    # 交通小计
    row += 1
    w(row, 2, "交通小计", font=Font(name="微软雅黑", bold=True, size=10), fill=subtotal_fill)
    w(row, 3, trip_t if trip_t else "-", font=Font(name="微软雅黑", bold=True, size=10),
      fill=subtotal_fill, align=right_align, fmt=amount_fmt if trip_t else None)
    totals["交通"] += trip_t

    # ── 2. 酒店 ──
    trip_h = 0
    for water, invoice in data["hotel_pairs"]:
        row += 1

        ref_file = water or invoice
        city_m = re.search(r"(\S+)住宿", ref_file["filename"])
        city = city_m.group(1) if city_m else ""

        (start, end) = parse_date_range(ref_file["filename"])
        date_range = f"{start[0]:02d}/{start[1]:02d}-{end[0]:02d}/{end[1]:02d}"

        amt = None
        if invoice:
            amt = extract_pdf_amount(invoice["fullpath"])
        if amt is None and water:
            amt = extract_pdf_amount(water["fullpath"])

        if amt:
            trip_h += amt

        has_invoice = "✓" if invoice else "✗ 缺发票"

        w(row, 1, "酒店住宿")
        w(row, 2, date_range, align=center_align)
        w(row, 3, amt if amt else "?", align=right_align, fmt=amount_fmt if amt else None)
        w(row, 4, city)
        w(row, 5, has_invoice)

    row += 1
    w(row, 2, "酒店小计", font=Font(name="微软雅黑", bold=True, size=10), fill=subtotal_fill)
    w(row, 3, trip_h if trip_h else "-", font=Font(name="微软雅黑", bold=True, size=10),
      fill=subtotal_fill, align=right_align, fmt=amount_fmt if trip_h else None)
    totals["酒店"] += trip_h

    # ── 3. 滴滴出行 ──
    didi_form_paths = [f["fullpath"] for _, f, _ in data["didi_paired"] if f]
    didi_inv_paths = [i["fullpath"] for _, _, i in data["didi_paired"] if i]

    didi_total = 0
    for ip in didi_inv_paths:
        amt = extract_pdf_amount(ip)
        if amt:
            didi_total += amt

    # 最后一张交通票的到达城市和日期
    last_arrival = ""
    last_return_date = ""
    for t in reversed(transport):
        route_m = re.search(r"(\S+)-(\S+)", t["filename"])
        if route_m:
            last_arrival = route_m.group(2)
            m = re.search(r"(\d{4})", t["filename"])
            if m:
                d = m.group(1)
                last_return_date = f"{d[:2]}/{d[2:]}"
            break

    row += 1
    w(row, 1, "滴滴出行")
    w(row, 2, f"共{len(didi_form_paths)}段行程", align=center_align)
    w(row, 3, didi_total if didi_total > 0 else "?", align=right_align,
      fmt=amount_fmt if didi_total > 0 else None)
    w(row, 4, f"返回{last_arrival}: {last_return_date}" if last_return_date else "")
    w(row, 5, f"行程单{len(didi_form_paths)}份, 发票{len(didi_inv_paths)}份")
    if didi_total > 0:
        totals["滴滴"] += didi_total

    # ── 4. 餐饮 ──
    trip_dining = 0
    if data["dining"]:
        for d in sorted(data["dining"], key=lambda x: parse_date_range(x["filename"])[0]):
            # 餐饮属于本地时才列出，行程中的餐饮通常是和本地混在一起的
            # 这里跳过，统一在本地部分处理
            pass

    # ── 行程合计 ──
    row += 1
    trip_grand = trip_t + trip_h + (didi_total if didi_total > 0 else 0)
    merge_and_write(row, 1, 5,
                    f"本段合计: ¥{trip_grand:,.2f}" if trip_grand > 0 else "本段合计: -",
                    font=Font(name="微软雅黑", bold=True, size=10), fill=trip_total_fill)
    row += 1  # 空行

# ── 本地费用 ──
local = classify_local(BASE_DIR)
local_t = {"滴滴": 0, "餐饮": 0, "通信": 0}

if local:
    merge_and_write(row, 1, 5, "【本地费用】",
                    font=Font(name="微软雅黑", bold=True, size=12), fill=cat_fill)
    row += 1

    # 通信
    if local.get("telecom"):
        for t in local["telecom"]:
            row += 1
            name = t["filename"]
            month_m = re.search(r"(\d{2})", name)
            month = month_m.group(1) + "月" if month_m else ""
            amt_m = re.search(r"(\d+\.?\d*)\.pdf", name)
            amt = float(amt_m.group(1)) if amt_m else None
            if amt:
                local_t["通信"] += amt

            w(row, 1, "通信费")
            w(row, 2, month, align=center_align)
            w(row, 3, amt if amt else "-", align=right_align, fmt=amount_fmt if amt else None)
            w(row, 4, "")
            w(row, 5, name)

        row += 1
        w(row, 2, "通信小计", font=Font(name="微软雅黑", bold=True, size=10), fill=subtotal_fill)
        w(row, 3, local_t["通信"], font=Font(name="微软雅黑", bold=True, size=10),
          fill=subtotal_fill, align=right_align, fmt=amount_fmt)
        totals["通信"] += local_t["通信"]

    # 滴滴(本地)
    if local.get("didi_paired"):
        f_paths = [f["fullpath"] for _, f, _ in local["didi_paired"] if f]
        i_paths = [i["fullpath"] for _, _, i in local["didi_paired"] if i]
        d_total = 0
        for ip in i_paths:
            amt = extract_pdf_amount(ip)
            if amt:
                d_total += amt
        if d_total > 0:
            local_t["滴滴"] += d_total
            totals["滴滴"] += d_total

        row += 1
        w(row, 1, "滴滴出行(本地)")
        w(row, 2, f"共{len(f_paths)}段", align=center_align)
        w(row, 3, d_total if d_total > 0 else "?", align=right_align,
          fmt=amount_fmt if d_total > 0 else None)
        w(row, 4, "")
        w(row, 5, "")

    # 餐饮(本地)
    if local.get("dining"):
        for d in local["dining"]:
            row += 1
            name = d["filename"]
            date_m = re.search(r"(\d{4})", name)
            mm_dd = date_m.group(1) if date_m else ""
            date_str = f"{mm_dd[:2]}/{mm_dd[2:]}" if mm_dd else ""
            amt_m = re.search(r"(\d+\.?\d*)\.pdf", name)
            amt = float(amt_m.group(1)) if amt_m else None
            if amt:
                local_t["餐饮"] += amt

            w(row, 1, "餐饮")
            w(row, 2, date_str, align=center_align)
            w(row, 3, amt if amt else "-", align=right_align, fmt=amount_fmt if amt else None)
            w(row, 4, "")
            w(row, 5, name)

        totals["餐饮"] += local_t["餐饮"]

    # 本地合计
    row += 1
    local_total = local_t["通信"] + local_t["滴滴"] + local_t["餐饮"]
    merge_and_write(row, 1, 5,
                    f"本地合计: ¥{local_total:,.2f}" if local_total > 0 else "本地合计: -",
                    font=Font(name="微软雅黑", bold=True, size=10), fill=trip_total_fill)

# ── 总计 ──
row += 2
grand_total = sum(totals.values())
grand_font_w = Font(name="微软雅黑", bold=True, size=13, color="FFFFFF")
detail = f"交通:{totals['交通']:,.0f}  酒店:{totals['酒店']:,.0f}  滴滴:{totals['滴滴']:,.0f}"
if totals['餐饮'] > 0:
    detail += f"  餐饮:{totals['餐饮']:,.0f}"
if totals['通信'] > 0:
    detail += f"  通信:{totals['通信']:,.0f}"

merge_and_write(row, 1, 5,
                f"报销总计: ¥{grand_total:,.2f}  ({detail})",
                font=grand_font_w, fill=grand_fill)

# 打印设置
ws.page_setup.orientation = "landscape"
ws.auto_filter.ref = f"A1:E{row}"

wb.save(str(OUTPUT))
print(f"Excel 已生成: {OUTPUT}")
print(f"交通: ¥{totals['交通']:,.2f}  酒店: ¥{totals['酒店']:,.2f}  滴滴: ¥{totals['滴滴']:,.2f}  餐饮: ¥{totals['餐饮']:,.2f}  通信: ¥{totals['通信']:,.2f}")
print(f"总计: ¥{grand_total:,.2f}")

# 检查缺失金额
missing = []
for trip_dir in trip_dirs:
    data = classify_and_sort(trip_dir, BASE_DIR)
    if data:
        for _, invoice in data["hotel_pairs"]:
            if invoice:
                amt = extract_pdf_amount(invoice["fullpath"])
                if amt is None:
                    missing.append(f"[酒店] {invoice['filename']}")

if missing:
    print(f"\n⚠️  无法提取金额的票据 ({len(missing)}项):")
    for m in missing:
        print(f"  {m}")
